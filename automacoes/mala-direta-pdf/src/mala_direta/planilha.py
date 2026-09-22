"""Leitura e validação da planilha de destinatários (.xlsx, .xlsm, .csv, .tsv)."""

from __future__ import annotations

import csv
import datetime as dt
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from mala_direta.config import ConfigPlanilha
from mala_direta.erros import PlanilhaInvalidaError

EMAIL = re.compile(r"^[^@\s]+@[^@\s.]+\.[^@\s]+$")
PLANILHAS_EXCEL = {".xlsx", ".xlsm"}
SEPARADORES = {".csv": ",", ".tsv": "\t"}


@dataclass(frozen=True)
class Destinatario:
    """Uma linha validada da planilha, com os valores já em texto."""

    linha: int
    email: str
    valores: dict[str, str]

    def valor(self, coluna: str) -> str:
        """Busca uma coluna ignorando maiúsculas, acentos e pontuação do título."""
        procurada = chave_de_coluna(coluna)
        for titulo, valor in self.valores.items():
            if chave_de_coluna(titulo) == procurada:
                return valor
        raise KeyError(coluna)


def ler_destinatarios(config: ConfigPlanilha, colunas_exigidas: list[str]) -> list[Destinatario]:
    """Lê a planilha e devolve as linhas válidas, ou falha listando todos os problemas."""
    cabecalho, linhas = _ler_tabela(config.arquivo, config.aba)
    indice = _mapear_colunas(cabecalho)

    exigidas = [config.coluna_email, *colunas_exigidas]
    ausentes = [
        coluna for coluna in dict.fromkeys(exigidas) if chave_de_coluna(coluna) not in indice
    ]
    if ausentes:
        raise PlanilhaInvalidaError(
            f"Colunas ausentes na planilha {config.arquivo.name}: {', '.join(ausentes)}.\n"
            f"Colunas encontradas: {', '.join(cabecalho) or '(nenhuma)'}"
        )

    destinatarios: list[Destinatario] = []
    problemas: list[str] = []

    for numero, linha in enumerate(linhas, start=2):
        valores = {
            titulo: _formatar(linha[posicao]) if posicao < len(linha) else ""
            for posicao, titulo in enumerate(cabecalho)
            if titulo
        }
        if not any(valores.values()):
            continue

        email = valores[cabecalho[indice[chave_de_coluna(config.coluna_email)]]].strip()
        if not EMAIL.match(email):
            problemas.append(f"linha {numero}: e-mail inválido ou vazio ({email or 'vazio'})")
            continue

        destinatarios.append(Destinatario(linha=numero, email=email, valores=valores))

    if problemas:
        raise PlanilhaInvalidaError(
            "Planilha com linhas inválidas:\n" + "\n".join(f"  - {p}" for p in problemas)
        )
    if not destinatarios:
        raise PlanilhaInvalidaError(f"Nenhuma linha preenchida em {config.arquivo.name}.")

    return destinatarios


def emails_repetidos(destinatarios: list[Destinatario]) -> list[str]:
    """E-mails que aparecem em mais de uma linha, para a CLI avisar antes de enviar."""
    contagem = Counter(destinatario.email.lower() for destinatario in destinatarios)
    return sorted(email for email, vezes in contagem.items() if vezes > 1)


def _ler_tabela(arquivo: Path, aba: str | None) -> tuple[list[str], list[list[Any]]]:
    if not arquivo.is_file():
        raise PlanilhaInvalidaError(f"Planilha não encontrada: {arquivo}")

    extensao = arquivo.suffix.lower()
    if extensao in PLANILHAS_EXCEL:
        return _ler_excel(arquivo, aba)
    if extensao in SEPARADORES:
        return _ler_texto(arquivo, SEPARADORES[extensao])
    raise PlanilhaInvalidaError(
        f"Formato não suportado: {extensao or arquivo.name}. Use .xlsx, .xlsm, .csv ou .tsv."
    )


def _ler_excel(arquivo: Path, aba: str | None) -> tuple[list[str], list[list[Any]]]:
    from openpyxl import load_workbook

    planilha = load_workbook(arquivo, data_only=True, read_only=True)
    try:
        if aba is not None and aba not in planilha.sheetnames:
            raise PlanilhaInvalidaError(
                f"Aba {aba!r} não existe em {arquivo.name}. "
                f"Abas disponíveis: {', '.join(planilha.sheetnames)}"
            )
        folha = planilha[aba] if aba else planilha[planilha.sheetnames[0]]
        linhas = [list(linha) for linha in folha.iter_rows(values_only=True)]
    finally:
        planilha.close()

    if not linhas:
        raise PlanilhaInvalidaError(f"Planilha vazia: {arquivo.name}")
    return [_formatar(celula) for celula in linhas[0]], linhas[1:]


def _ler_texto(arquivo: Path, separador: str) -> tuple[list[str], list[list[Any]]]:
    with arquivo.open(encoding="utf-8-sig", newline="") as texto:
        linhas = [linha for linha in csv.reader(texto, delimiter=separador)]
    if not linhas:
        raise PlanilhaInvalidaError(f"Planilha vazia: {arquivo.name}")
    return [celula.strip() for celula in linhas[0]], linhas[1:]


def _mapear_colunas(cabecalho: list[str]) -> dict[str, int]:
    """Índice tolerante a maiúsculas, acentos e pontuação: 'E-mail' encontra 'email'."""
    indice: dict[str, int] = {}
    for posicao, titulo in enumerate(cabecalho):
        chave = chave_de_coluna(titulo)
        if chave and chave not in indice:
            indice[chave] = posicao
    return indice


def chave_de_coluna(titulo: str) -> str:
    """Forma canônica de um nome de coluna: sem acento, minúsculo, sem pontuação."""
    sem_acento = unicodedata.normalize("NFKD", titulo).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", sem_acento.lower())


def _formatar(celula: Any) -> str:
    """Converte a célula em texto do jeito que uma pessoa espera ver no PDF."""
    if celula is None:
        return ""
    if isinstance(celula, bool):
        return "sim" if celula else "não"
    if isinstance(celula, dt.datetime):
        return celula.strftime("%d/%m/%Y")
    if isinstance(celula, dt.date):
        return celula.strftime("%d/%m/%Y")
    if isinstance(celula, float) and celula.is_integer():
        return str(int(celula))
    return str(celula).strip()
